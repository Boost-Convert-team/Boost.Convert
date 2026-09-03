import sys
import tempfile
import types
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

import fitz
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from openpyxl import Workbook

BACKEND_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_ROOT))

from Blueprints.services.convertions_services.documents import (
    docx_to_pdf_service,
    office_pdf_fallback,
    xlsx_to_pdf_service,
)


class OfficePdfFallbackTests(unittest.TestCase):
    def test_fixed_table_html_preserves_word_layout_properties(self) -> None:
        document = Document()
        section = document.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11.69)
        section.page_height = Inches(8.27)
        section.left_margin = Inches(0.3)

        table = document.add_table(rows=3, cols=3)
        for grid_column, width in zip(
            table._tbl.tblGrid.gridCol_lst,
            ("2500", "4000", "6000"),
        ):
            grid_column.set(qn("w:w"), width)

        merged_cell = table.cell(0, 0).merge(table.cell(0, 1))
        merged_cell.text = "ÁREA MESCLADA"
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:fill"), "BFBFBF")
        merged_cell._tc.get_or_add_tcPr().append(shading)

        break_run = table.cell(2, 0).paragraphs[0].add_run("SEGUNDA PÁGINA")
        break_run._r.insert(0, OxmlElement("w:lastRenderedPageBreak"))

        content = office_pdf_fallback.build_fixed_table_html(document)

        self.assertIn("@page{size:841.7pt 595.45pt;margin:0;}", content)
        self.assertIn('<col style="width:125pt">', content)
        self.assertIn('<col style="width:200pt">', content)
        self.assertIn('<col style="width:300pt">', content)
        self.assertIn('colspan="2"', content)
        self.assertIn("background:#BFBFBF", content)
        self.assertIn("page-break-before:always", content)
        self.assertEqual(content.count("ÁREA MESCLADA"), 1)
        self.assertEqual(content.count("SEGUNDA PÁGINA"), 1)

    def test_docx_to_pdf_converts_without_word_or_libreoffice(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "relatorio.docx"
            output = root / "relatorio.pdf"

            document = Document()
            document.add_heading("Relatorio fallback DOCX", level=1)
            document.add_paragraph("Texto convertido sem Word e sem LibreOffice.")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Campo"
            table.cell(0, 1).text = "Resultado"
            table.cell(1, 0).text = "Motor"
            table.cell(1, 1).text = "Python"
            document.save(source)

            docx2pdf = types.SimpleNamespace(
                convert=Mock(side_effect=RuntimeError("Word unavailable"))
            )
            with (
                patch.dict(sys.modules, {"docx2pdf": docx2pdf}),
                patch.object(
                    docx_to_pdf_service,
                    "run_libreoffice_conversion",
                    side_effect=RuntimeError("LibreOffice unavailable"),
                ),
            ):
                docx_to_pdf_service.convert_docx_pdf(source, output)

            text = extract_pdf_text(output)
            self.assertGreater(output.stat().st_size, 0)
            self.assertIn("Relatorio fallback DOCX", text)
            self.assertIn("Texto convertido sem Word", text)
            self.assertIn("Python", text)

    def test_xlsx_to_pdf_converts_without_libreoffice(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "planilha.xlsx"
            output = root / "planilha.pdf"

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Resumo"
            sheet.append(["Produto", "Quantidade", "Total"])
            sheet.append(["Plano Pro", 3, 149.70])
            sheet.append(["Plano Free", 7, 0])
            workbook.save(source)

            with patch.object(
                xlsx_to_pdf_service,
                "run_libreoffice_conversion",
                side_effect=RuntimeError("LibreOffice unavailable"),
            ):
                xlsx_to_pdf_service.convert_excel_pdf(source, output)

            text = extract_pdf_text(output)
            self.assertGreater(output.stat().st_size, 0)
            self.assertIn("Resumo", text)
            self.assertIn("Produto", text)
            self.assertIn("Plano Pro", text)
            self.assertIn("149.7", text)


def extract_pdf_text(path: Path) -> str:
    with fitz.open(path) as document:
        return "\n".join(page.get_text("text") for page in document)


if __name__ == "__main__":
    unittest.main()
