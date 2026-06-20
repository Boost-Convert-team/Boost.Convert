import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

import fitz
from docx import Document


BACKEND_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_ROOT))

from Blueprints.services.ai import document_text_extractor


class DocumentTextExtractorTests(unittest.TestCase):
    def test_extract_document_text_dispatches_and_normalizes(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "notas.txt"
            path.write_text("texto   claro\n\n\nfim", encoding="utf-8")

            text = document_text_extractor.extract_document_text(path, "txt")

        self.assertEqual(text, "texto claro\n\nfim")

    def test_dispatch_document_extraction_uses_extension_parser(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "notas.md"
            path.write_text("markdown", encoding="utf-8")

            text = document_text_extractor.dispatch_document_extraction(path, "md")

        self.assertEqual(text, "markdown")

    def test_normalize_document_text_collapses_whitespace(self) -> None:
        text = document_text_extractor.normalize_document_text("a   b\n\n\nc")

        self.assertEqual(text, "a b\n\nc")

    def test_extract_plain_text_reads_utf8(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "notas.txt"
            path.write_text("conteudo claro", encoding="utf-8")

            text = document_text_extractor.extract_plain_text(path)

        self.assertEqual(text, "conteudo claro")

    def test_extract_docx_text_reads_paragraphs(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "notas.docx"
            document = Document()
            document.add_paragraph("Primeiro")
            document.add_paragraph("Segundo")
            document.save(path)

            text = document_text_extractor.extract_docx_text(path)

        self.assertEqual(text, "Primeiro\nSegundo")

    def test_extract_pdf_text_reads_page_text(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "notas.pdf"
            document = fitz.open()
            page = document.new_page()
            page.insert_text((72, 72), "Texto do PDF")
            document.save(path)
            document.close()

            text = document_text_extractor.extract_pdf_text(path)

        self.assertIn("Texto do PDF", text)

    def test_extract_odt_text_reads_content_xml(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "notas.odt"
            xml = "<office><body><p>Texto ODT</p></body></office>"
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("content.xml", xml)

            text = document_text_extractor.extract_odt_text(path)

        self.assertIn("Texto ODT", text)

    def test_extract_rtf_text_removes_control_words(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "notas.rtf"
            path.write_text(r"{\rtf1\ansi Texto RTF}", encoding="utf-8")

            text = document_text_extractor.extract_rtf_text(path)

        self.assertIn("Texto RTF", text)
        self.assertNotIn("rtf1", text)

    def test_extract_doc_text_uses_libreoffice_wrapper(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "legado.doc"

            with patch.object(
                document_text_extractor,
                "extract_libreoffice_text",
                return_value="texto legado",
            ) as extractor:
                text = document_text_extractor.extract_doc_text(path)

        self.assertEqual(text, "texto legado")
        extractor.assert_called_once_with(path)

    def test_validate_document_extension_rejects_unknown_extension(self) -> None:
        with self.assertRaisesRegex(ValueError, "Arquivo invalido"):
            document_text_extractor.validate_document_extension("exe")

    def test_validate_document_size_rejects_large_file(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "grande.txt"
            limit = document_text_extractor.DOCUMENT_ANALYZER_MAX_UPLOAD_BYTES + 1
            path.write_bytes(b"0" * limit)

            with self.assertRaisesRegex(ValueError, "Documento muito grande"):
                document_text_extractor.validate_document_size(path)

    def test_validate_extracted_document_text_rejects_empty_text(self) -> None:
        with self.assertRaisesRegex(ValueError, "Texto nao extraido"):
            document_text_extractor.validate_extracted_document_text("   ")


if __name__ == "__main__":
    unittest.main()
