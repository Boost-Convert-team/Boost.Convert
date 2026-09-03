import html
import logging
import math
import os
import re
import tempfile
from collections.abc import Iterable, Sequence
from datetime import date, datetime, time
from pathlib import Path

import fitz
from docx import Document
from docx.oxml.ns import qn
from docx.table import _Cell
from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils import get_column_letter

PathLike = str | os.PathLike[str]

A4 = fitz.paper_rect("a4")
A4_LANDSCAPE = fitz.Rect(0, 0, A4.height, A4.width)
MARGIN = 42
TEXT_COLOR = (0.12, 0.14, 0.18)
BORDER_COLOR = (0.78, 0.80, 0.84)
HEADER_FILL = (0.94, 0.95, 0.97)

logger = logging.getLogger(__name__)


class PdfWriter:
    def __init__(self, document: fitz.Document, page_rect: fitz.Rect = A4):
        self.document = document
        self.page_rect = page_rect
        self.page = self.document.new_page(width=page_rect.width, height=page_rect.height)
        self.y = MARGIN

    @property
    def usable_width(self) -> float:
        return self.page_rect.width - (MARGIN * 2)

    @property
    def bottom(self) -> float:
        return self.page_rect.height - MARGIN

    def ensure_space(self, height: float) -> None:
        if self.y + height <= self.bottom:
            return
        self.page = self.document.new_page(width=self.page_rect.width, height=self.page_rect.height)
        self.y = MARGIN

    def draw_heading(self, text: str, size: float = 18) -> None:
        height = max(32, estimate_text_height(text, size, self.usable_width))
        self.ensure_space(height)
        self.page.insert_textbox(
            fitz.Rect(MARGIN, self.y, self.page_rect.width - MARGIN, self.y + height),
            text,
            fontsize=size,
            fontname="hebo",
            color=TEXT_COLOR,
        )
        self.y += height + 8

    def draw_paragraph(self, text: str, size: float = 11) -> None:
        text = normalize_text(text)
        if not text:
            self.y += 8
            return

        height = max(28, estimate_text_height(text, size, self.usable_width) + 8)
        self.ensure_space(height)
        self.page.insert_textbox(
            fitz.Rect(MARGIN, self.y, self.page_rect.width - MARGIN, self.y + height),
            text,
            fontsize=size,
            fontname="helv",
            color=TEXT_COLOR,
        )
        self.y += height + 6

    def draw_table(self, rows: Sequence[Sequence[str]], header: bool = False) -> None:
        rows = [[normalize_text(cell) for cell in row] for row in rows if any(normalize_text(cell) for cell in row)]
        if not rows:
            return

        column_count = max(len(row) for row in rows)
        column_width = self.usable_width / max(1, column_count)

        self.y += 4
        for row_index, row in enumerate(rows):
            row_values = list(row) + [""] * (column_count - len(row))
            row_height = max(
                24,
                max(estimate_text_height(value, 9.2, column_width - 10) for value in row_values) + 16,
            )
            self.ensure_space(row_height)
            x = MARGIN
            fill = HEADER_FILL if header and row_index == 0 else None
            for value in row_values:
                rect = fitz.Rect(x, self.y, x + column_width, self.y + row_height)
                self.page.draw_rect(rect, color=BORDER_COLOR, fill=fill, width=0.45)
                self.page.insert_textbox(
                    rect + (5, 5, -5, -4),
                    value,
                    fontsize=9.2,
                    fontname="helv",
                    color=TEXT_COLOR,
                )
                x += column_width
            self.y += row_height
        self.y += 10


def convert_docx_to_pdf_fallback(input_path: PathLike, output_path: PathLike) -> None:
    source = Document(input_path)

    if _is_fixed_table_document(source):
        try:
            with tempfile.TemporaryDirectory() as temp_dir:
                rendered_path = Path(temp_dir) / "fixed-table.pdf"
                render_fixed_table_pdf_with_playwright(source, rendered_path)
                if rendered_path.exists() and rendered_path.stat().st_size > 0:
                    final_path = Path(output_path)
                    final_path.parent.mkdir(parents=True, exist_ok=True)
                    os.replace(rendered_path, final_path)
                    return
        except Exception:
            logger.debug("Fixed-layout DOCX rendering was unavailable.", exc_info=True)

    pdf = fitz.open()
    writer = PdfWriter(pdf)

    if source.core_properties.title:
        writer.draw_heading(source.core_properties.title)

    for block in iter_docx_blocks(source):
        if isinstance(block, list):
            writer.draw_table(block, header=True)
            continue

        text, style_name = block
        if style_name.startswith(("heading 1", "titulo 1")):
            writer.draw_heading(text, 18)
        elif style_name.startswith(("heading", "titulo")):
            writer.draw_heading(text, 14)
        else:
            writer.draw_paragraph(text)

    save_pdf(pdf, output_path)


def _is_fixed_table_document(document) -> bool:
    if len(document.sections) != 1 or len(document.tables) != 1:
        return False

    section = document.sections[0]
    header_footer_parts = (
        section.header,
        section.first_page_header,
        section.even_page_header,
        section.footer,
        section.first_page_footer,
        section.even_page_footer,
    )
    if any(
        paragraph.text.strip()
        for part in header_footer_parts
        for paragraph in part.paragraphs
    ):
        return False

    table_seen = False
    paragraph_map = {
        paragraph._element: paragraph for paragraph in document.paragraphs
    }
    for child in document.element.body.iterchildren():
        if child is document.tables[0]._element:
            table_seen = True
        elif child in paragraph_map and paragraph_map[child].text.strip() and not table_seen:
            return False
    return table_seen


def render_fixed_table_pdf_with_playwright(document, output_path: PathLike) -> None:
    from playwright.sync_api import sync_playwright

    content = build_fixed_table_html(document)
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(args=["--no-sandbox"])
        try:
            context = browser.new_context(java_script_enabled=False)
            context.route("http://*/*", lambda route: route.abort())
            context.route("https://*/*", lambda route: route.abort())
            page = context.new_page()
            page.set_content(content, wait_until="load", timeout=15000)
            page.emulate_media(media="print")
            page.pdf(
                path=os.fspath(output_path),
                print_background=True,
                prefer_css_page_size=True,
                margin={"top": "0", "right": "0", "bottom": "0", "left": "0"},
            )
            context.close()
        finally:
            browser.close()


def build_fixed_table_html(document) -> str:
    section = document.sections[0]
    page_width = _emu_to_points(section.page_width)
    page_height = _emu_to_points(section.page_height)
    left_margin = _emu_to_points(section.left_margin)
    table = document.tables[0]
    column_widths = [
        _safe_number(column.get(qn("w:w")), 1) / 20
        for column in table._tbl.tblGrid.gridCol_lst
    ]
    if not column_widths:
        raise ValueError("Tabela DOCX sem grade de colunas.")

    table_width = sum(column_widths)
    table_y = _safe_number(
        _word_attribute(table._tbl.tblPr, "w:tblpPr", "w:tblpY", "0"),
        0,
    ) / 20
    rows = _map_physical_table_cells(table)

    parts = [
        "<!doctype html><html><head><meta charset=\"utf-8\"><style>",
        f"@page{{size:{page_width:g}pt {page_height:g}pt;margin:0;}}",
        "*{box-sizing:border-box;-webkit-print-color-adjust:exact;print-color-adjust:exact}",
        "html,body{margin:0;padding:0;background:#fff;color:#000}",
        (
            f"table{{border-collapse:collapse;table-layout:fixed;width:{table_width:g}pt;"
            f"margin-left:{left_margin:g}pt;margin-top:{table_y:g}pt}}"
        ),
        "tr{break-inside:avoid;page-break-inside:avoid}",
        "td{border:.5pt solid #000;padding:1.5pt 3pt;overflow:hidden}",
        "p{padding:0}",
        "</style></head><body><table><colgroup>",
    ]
    parts.extend(f'<col style="width:{width:g}pt">' for width in column_widths)
    parts.append("</colgroup><tbody>")

    for row_index, mapped_cells in enumerate(rows):
        row = table.rows[row_index]
        row_height = _emu_to_points(row.height) if row.height else 0
        break_before = (
            "break-before:page;page-break-before:always;"
            if row._tr.xpath(".//w:lastRenderedPageBreak")
            else ""
        )
        parts.append(f'<tr style="{break_before}height:{row_height:g}pt">')

        for mapped_cell in mapped_cells:
            if mapped_cell["merge"] == "continue":
                continue

            table_cell = mapped_cell["element"]
            fill = _safe_hex_color(
                _word_attribute(table_cell.tcPr, "w:shd", "w:fill", "FFFFFF"),
                "FFFFFF",
            )
            vertical_alignment = _safe_alignment(
                _word_attribute(table_cell.tcPr, "w:vAlign", "w:val", "top"),
                vertical=True,
            )
            cell = _Cell(table_cell, table)
            cell_content = "".join(
                _paragraph_to_html(paragraph) for paragraph in cell.paragraphs
            )
            parts.append(
                f'<td colspan="{mapped_cell["span"]}" '
                f'rowspan="{mapped_cell.get("rowspan", 1)}" '
                f'style="background:#{fill};vertical-align:{vertical_alignment}">'
                f"{cell_content}</td>"
            )
        parts.append("</tr>")

    parts.append("</tbody></table>")
    if document.paragraphs:
        parts.append(
            f'<div style="width:{table_width:g}pt;margin-left:{left_margin:g}pt">'
        )
        parts.extend(
            _paragraph_to_html(paragraph)
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        )
        parts.append("</div>")
    parts.append("</body></html>")
    return "".join(parts)


def _map_physical_table_cells(table) -> list[list[dict[str, object]]]:
    rows: list[list[dict[str, object]]] = []
    for row in table.rows:
        column_index = 0
        mapped_cells = []
        for table_cell in row._tr.tc_lst:
            span = max(
                1,
                int(
                    _safe_number(
                        _word_attribute(table_cell.tcPr, "w:gridSpan", "w:val", "1"),
                        1,
                    )
                ),
            )
            vertical_merge = table_cell.tcPr.find(qn("w:vMerge"))
            merge_state = (
                None
                if vertical_merge is None
                else (vertical_merge.get(qn("w:val")) or "continue")
            )
            mapped_cells.append(
                {
                    "element": table_cell,
                    "column": column_index,
                    "span": span,
                    "merge": merge_state,
                }
            )
            column_index += span
        rows.append(mapped_cells)

    for row_index, mapped_cells in enumerate(rows):
        for mapped_cell in mapped_cells:
            if mapped_cell["merge"] != "restart":
                continue
            rowspan = 1
            for next_row in rows[row_index + 1 :]:
                continuation = next(
                    (
                        cell
                        for cell in next_row
                        if cell["column"] == mapped_cell["column"]
                    ),
                    None,
                )
                if not continuation or continuation["merge"] != "continue":
                    break
                rowspan += 1
            mapped_cell["rowspan"] = rowspan

    return rows


def _paragraph_to_html(paragraph) -> str:
    paragraph_properties = paragraph._p.pPr
    alignment = _safe_alignment(
        _word_attribute(paragraph_properties, "w:jc", "w:val", "left")
    )
    before = _safe_number(
        _word_attribute(paragraph_properties, "w:spacing", "w:before", "0"),
        0,
    ) / 20
    after = _safe_number(
        _word_attribute(paragraph_properties, "w:spacing", "w:after", "0"),
        0,
    ) / 20
    content = "".join(_run_to_html(run, paragraph) for run in paragraph.runs)
    return (
        f'<p style="text-align:{alignment};margin:{before:g}pt 0 {after:g}pt;'
        f'line-height:1.05">{content or "&nbsp;"}</p>'
    )


def _run_to_html(run, paragraph) -> str:
    if not run.text:
        return ""

    run_properties = run._r.rPr
    paragraph_properties = paragraph._p.pPr
    paragraph_run_properties = (
        paragraph_properties.find(qn("w:rPr"))
        if paragraph_properties is not None
        else None
    )
    font_name = (
        _word_attribute(run_properties, "w:rFonts", "w:ascii")
        or _word_attribute(paragraph_run_properties, "w:rFonts", "w:ascii")
        or "Arial"
    )
    font_name = re.sub(r"[^A-Za-z0-9 _-]", "", font_name)[:80] or "Arial"
    size = _safe_number(
        _word_attribute(run_properties, "w:sz", "w:val")
        or _word_attribute(paragraph_run_properties, "w:sz", "w:val")
        or "20",
        20,
    ) / 2
    color = _safe_hex_color(
        _word_attribute(run_properties, "w:color", "w:val")
        or _word_attribute(paragraph_run_properties, "w:color", "w:val")
        or "000000",
        "000000",
    )
    styles = [
        f"font-family:{html.escape(font_name, quote=True)},Arial,sans-serif",
        f"font-size:{size:g}pt",
        f"color:#{color}",
    ]
    if (
        _word_property_enabled(run_properties, "w:b")
        or _word_property_enabled(paragraph_run_properties, "w:b")
        or font_name.lower() == "arial black"
    ):
        styles.append("font-weight:700")
    if _word_property_enabled(run_properties, "w:i") or _word_property_enabled(
        paragraph_run_properties, "w:i"
    ):
        styles.append("font-style:italic")
    underline = _word_attribute(run_properties, "w:u", "w:val") or _word_attribute(
        paragraph_run_properties, "w:u", "w:val"
    )
    if underline and underline != "none":
        styles.append("text-decoration:underline")

    text = (
        html.escape(run.text)
        .replace("\t", " ")
        .replace("\r\n", "<br>")
        .replace("\r", "<br>")
        .replace("\n", "<br>")
    )
    return f'<span style="{";".join(styles)}">{text}</span>'


def _word_attribute(element, child_name: str, attribute_name: str, default=None):
    child = element.find(qn(child_name)) if element is not None else None
    value = child.get(qn(attribute_name)) if child is not None else None
    return default if value is None else value


def _word_property_enabled(element, child_name: str) -> bool:
    child = element.find(qn(child_name)) if element is not None else None
    if child is None:
        return False
    return child.get(qn("w:val"), "1").lower() not in {
        "0",
        "false",
        "off",
        "none",
    }


def _safe_hex_color(value, default: str) -> str:
    value = str(value or "")
    if value.lower() in {"auto", "none"}:
        return default
    if re.fullmatch(r"[0-9A-Fa-f]{6}|[0-9A-Fa-f]{8}", value):
        return value[-6:].upper()
    return default


def _safe_alignment(value, vertical: bool = False) -> str:
    if vertical:
        return value if value in {"top", "center", "bottom"} else "top"
    return {
        "left": "left",
        "center": "center",
        "right": "right",
        "both": "justify",
        "distribute": "justify",
    }.get(value, "left")


def _safe_number(value, default: float) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _emu_to_points(value) -> float:
    return _safe_number(value, 0) / 12700


def convert_xlsx_to_pdf_fallback(input_path: PathLike, output_path: PathLike) -> None:
    workbook = load_workbook(input_path, data_only=True)
    pdf = fitz.open()

    for sheet in workbook.worksheets:
        if sheet.sheet_state != "visible":
            continue
        rows = get_sheet_rows(sheet)
        if not rows:
            continue

        writer = PdfWriter(pdf, A4_LANDSCAPE)
        writer.draw_heading(sheet.title, 16)
        draw_sheet_table(writer, sheet, rows)

    if pdf.page_count == 0:
        writer = PdfWriter(pdf, A4_LANDSCAPE)
        writer.draw_paragraph("Planilha sem dados visiveis.")

    save_pdf(pdf, output_path)


def iter_docx_blocks(document: Document) -> Iterable[tuple[str, str] | list[list[str]]]:
    body = document.element.body
    paragraph_map = {paragraph._element: paragraph for paragraph in document.paragraphs}
    table_map = {table._element: table for table in document.tables}

    for child in body.iterchildren():
        if child in paragraph_map:
            paragraph = paragraph_map[child]
            text = paragraph.text.strip()
            style_name = (paragraph.style.name if paragraph.style else "").strip().lower()
            if text:
                yield text, style_name
        elif child in table_map:
            table = table_map[child]
            yield [[cell.text.strip() for cell in row.cells] for row in table.rows]


def get_sheet_rows(sheet) -> list[list[Cell]]:
    rows = []
    for row in sheet.iter_rows():
        if any(cell.value not in (None, "") for cell in row):
            rows.append(list(row))

    while rows and not any(cell.value not in (None, "") for cell in rows[-1]):
        rows.pop()

    if not rows:
        return []

    last_column = max(
        index
        for row in rows
        for index, cell in enumerate(row)
        if cell.value not in (None, "")
    )
    return [row[: last_column + 1] for row in rows]


def draw_sheet_table(writer: PdfWriter, sheet, rows: Sequence[Sequence[Cell]]) -> None:
    widths = get_scaled_column_widths(sheet, len(rows[0]), writer.usable_width)
    for row_index, row in enumerate(rows):
        values = [format_cell_value(cell) for cell in row]
        heights = [
            estimate_text_height(value, 8.5, max(24, widths[index] - 8)) + 14
            for index, value in enumerate(values)
        ]
        row_height = max(22, min(78, max(heights) if heights else 22))
        writer.ensure_space(row_height)

        x = MARGIN
        for column_index, cell in enumerate(row):
            width = widths[column_index]
            rect = fitz.Rect(x, writer.y, x + width, writer.y + row_height)
            fill = get_cell_fill(cell, row_index)
            writer.page.draw_rect(rect, color=BORDER_COLOR, fill=fill, width=0.45)
            writer.page.insert_textbox(
                rect + (4, 4, -4, -3),
                values[column_index],
                fontsize=8.5,
                fontname="hebo" if row_index == 0 else "helv",
                color=TEXT_COLOR,
            )
            x += width
        writer.y += row_height


def get_scaled_column_widths(sheet, column_count: int, usable_width: float) -> list[float]:
    widths = []
    for index in range(1, column_count + 1):
        letter = get_column_letter(index)
        configured_width = sheet.column_dimensions[letter].width or 10
        widths.append(max(42, min(150, configured_width * 6.2)))

    total_width = sum(widths) or 1
    scale = usable_width / total_width
    return [width * scale for width in widths]


def format_cell_value(cell: Cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, float):
        if math.isfinite(value) and value.is_integer():
            return str(int(value))
        return f"{value:.2f}".rstrip("0").rstrip(".")
    return str(value)


def get_cell_fill(cell: Cell, row_index: int) -> tuple[float, float, float] | None:
    color = cell.fill.fgColor
    rgb = color.rgb if color and color.type == "rgb" else None
    if rgb and len(rgb) == 8 and rgb[:2] != "00":
        try:
            return tuple(int(rgb[index : index + 2], 16) / 255 for index in (2, 4, 6))
        except ValueError:
            return HEADER_FILL if row_index == 0 else None
    return HEADER_FILL if row_index == 0 else None


def estimate_text_height(text: str, font_size: float, width: float) -> float:
    if not text:
        return font_size * 1.35
    average_char_width = max(1, font_size * 0.48)
    chars_per_line = max(8, int(width / average_char_width))
    lines = sum(max(1, math.ceil(len(line) / chars_per_line)) for line in text.splitlines() or [text])
    return lines * font_size * 1.65


def normalize_text(value) -> str:
    return str(value or "").replace("\t", " ").strip()


def save_pdf(pdf: fitz.Document, output_path: PathLike) -> None:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    pdf.save(path)
    pdf.close()

    if not path.exists() or path.stat().st_size == 0:
        raise RuntimeError("Arquivo final nao foi criado.")
