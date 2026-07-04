import io
import os
import tempfile
from dataclasses import dataclass
from multiprocessing import Pool, cpu_count
from pathlib import Path
from typing import Union

import fitz
from pdf2docx import Converter


PathLike = Union[str, os.PathLike[str]]
DISABLE_PARALLEL_ENV = "PDF_TO_DOCX_DISABLE_PARALLEL"
MAX_WORKERS_ENV = "PDF_TO_DOCX_MAX_WORKERS"
MODE_ENV = "PDF_TO_DOCX_MODE"
VISUAL_DPI_ENV = "PDF_TO_DOCX_VISUAL_DPI"

DEFAULT_VISUAL_DPI = 100
MAX_PROFILE_SAMPLE_PAGES = 12
VISUAL_MODE = "visual"
EDITABLE_MODE = "editable"
AUTO_MODE = "auto"


@dataclass(frozen=True)
class PdfComplexityProfile:
    page_count: int
    average_images_per_page: float
    average_drawings_per_page: float

    @property
    def estimated_images(self) -> float:
        return self.average_images_per_page * self.page_count

    @property
    def estimated_drawings(self) -> float:
        return self.average_drawings_per_page * self.page_count


class ParallelPdfToDocxConverter(Converter):
    def _convert_with_multi_processing(self, docx_filename: str, start: int, end: int, **kwargs) -> None:
        workers = get_worker_count_from_settings(kwargs)
        if workers <= 1:
            sequential_settings = dict(kwargs, multi_processing=False, cpu_count=1)
            self.parse(start, end, None, **sequential_settings).make_docx(docx_filename, **sequential_settings)
            return

        temp_parent = get_temp_parent(docx_filename)
        with tempfile.TemporaryDirectory(prefix="pdf2docx-", dir=temp_parent) as temp_dir:
            temp_path = Path(temp_dir)
            vectors = [
                (
                    worker_index,
                    workers,
                    start,
                    end,
                    self.filename_pdf,
                    self.password,
                    kwargs,
                    str(temp_path / f"pages-{worker_index}.json"),
                )
                for worker_index in range(workers)
            ]

            with Pool(processes=workers) as pool:
                pool.map(Converter._parse_pages_per_cpu, vectors, 1)

            for worker_index in range(workers):
                filename = temp_path / f"pages-{worker_index}.json"
                if not filename.exists():
                    continue
                self.deserialize(str(filename))
                filename.unlink()

        self.make_docx(docx_filename, **kwargs)


def convert_pdf_docx(input_path: PathLike, output_path: PathLike) -> None:
    input_filename = os.fspath(input_path)
    output_filename = os.fspath(output_path)

    try:
        profile = inspect_pdf_complexity(input_filename)
        if should_use_visual_docx_conversion(profile):
            convert_pdf_to_visual_docx(input_filename, output_filename)
            return

        converter = ParallelPdfToDocxConverter(input_filename)
        try:
            converter.convert(output_filename, **build_pdf_to_docx_settings(profile.page_count))
        finally:
            converter.close()

    except Exception as exc:
        raise RuntimeError(
            f"Nao foi possivel preservar o layout do PDF em DOCX: {input_path}; esperado PDF compativel com pdf2docx."
        ) from exc


def inspect_pdf_complexity(input_path: PathLike) -> PdfComplexityProfile:
    with fitz.open(input_path) as document:
        page_count = len(document)
        if page_count == 0:
            return PdfComplexityProfile(0, 0, 0)

        sample_pages = get_profile_sample_pages(page_count)
        image_count = 0
        drawing_count = 0
        for page_index in sample_pages:
            page = document[page_index]
            image_count += len(page.get_images(full=True))
            drawing_count += len(page.get_drawings())

    sample_count = len(sample_pages)
    return PdfComplexityProfile(
        page_count=page_count,
        average_images_per_page=image_count / sample_count,
        average_drawings_per_page=drawing_count / sample_count,
    )


def get_profile_sample_pages(page_count: int) -> list[int]:
    if page_count <= MAX_PROFILE_SAMPLE_PAGES:
        return list(range(page_count))

    last_index = page_count - 1
    return sorted(
        {
            round(index * last_index / (MAX_PROFILE_SAMPLE_PAGES - 1))
            for index in range(MAX_PROFILE_SAMPLE_PAGES)
        }
    )


def should_use_visual_docx_conversion(profile: PdfComplexityProfile) -> bool:
    mode = os.getenv(MODE_ENV, AUTO_MODE).strip().lower()
    if mode == VISUAL_MODE:
        return True
    if mode == EDITABLE_MODE:
        return False

    if profile.page_count >= 80 and profile.average_images_per_page >= 25:
        return True

    if profile.page_count >= 40 and profile.estimated_drawings >= 5000:
        return True

    if profile.page_count >= 40 and profile.estimated_images >= 3000:
        return True

    return False


def convert_pdf_to_visual_docx(input_path: PathLike, output_path: PathLike, dpi: int | None = None) -> None:
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt

    render_dpi = dpi or get_int_env(VISUAL_DPI_ENV, DEFAULT_VISUAL_DPI)
    document = Document()
    configure_visual_docx_defaults(document)

    with fitz.open(input_path) as pdf:
        if len(pdf) == 0:
            raise RuntimeError("PDF sem paginas para converter.")

        configure_visual_docx_section(document.sections[0], pdf[0].rect)
        matrix = fitz.Matrix(render_dpi / 72, render_dpi / 72)

        for page_index, page in enumerate(pdf):
            paragraph = document.add_paragraph()
            if page_index > 0:
                paragraph.add_run().add_break(WD_BREAK.PAGE)

            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1

            pixmap = page.get_pixmap(matrix=matrix, alpha=False)
            image_stream = io.BytesIO(pixmap.tobytes("png"))
            paragraph.add_run().add_picture(
                image_stream,
                width=Pt(page.rect.width),
                height=Pt(page.rect.height),
            )

    document.save(output_path)


def configure_visual_docx_defaults(document) -> None:
    from docx.shared import Pt

    style = document.styles["Normal"]
    style.font.size = Pt(1)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1


def configure_visual_docx_section(section, page_rect) -> None:
    from docx.shared import Pt

    section.page_width = Pt(page_rect.width)
    section.page_height = Pt(page_rect.height)
    section.top_margin = Pt(0)
    section.bottom_margin = Pt(0)
    section.left_margin = Pt(0)
    section.right_margin = Pt(0)
    section.header_distance = Pt(0)
    section.footer_distance = Pt(0)


def build_pdf_to_docx_settings(page_count: int) -> dict[str, int | bool]:
    workers = get_parallel_worker_count(page_count)
    return {
        "multi_processing": workers > 1,
        "cpu_count": workers,
    }


def get_parallel_worker_count(page_count: int) -> int:
    if page_count <= 1 or get_bool_env(DISABLE_PARALLEL_ENV, False):
        return 1

    configured_limit = get_int_env(MAX_WORKERS_ENV, get_system_cpu_count())
    return max(1, min(page_count, configured_limit, get_system_cpu_count()))


def get_worker_count_from_settings(settings: dict) -> int:
    try:
        configured_workers = int(settings.get("cpu_count") or 0)
    except (TypeError, ValueError):
        configured_workers = 0
    if configured_workers <= 0:
        configured_workers = get_system_cpu_count()
    return max(1, min(configured_workers, get_system_cpu_count()))


def get_system_cpu_count() -> int:
    try:
        return max(1, cpu_count())
    except NotImplementedError:
        return 1


def get_temp_parent(docx_filename) -> str | None:
    if not isinstance(docx_filename, (str, bytes, os.PathLike)):
        return None

    parent = Path(os.fsdecode(docx_filename)).resolve().parent
    if parent.exists():
        return str(parent)
    return None


def get_int_env(name: str, default: int) -> int:
    try:
        return max(1, int(os.getenv(name, str(default))))
    except ValueError:
        return default


def get_bool_env(name: str, default: bool) -> bool:
    value = os.getenv(name)
    if value is None:
        return default
    return value.strip().lower() in {"1", "true", "yes", "on"}
