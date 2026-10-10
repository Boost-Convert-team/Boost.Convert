import os
import shutil
import subprocess
import tempfile
import wave
from pathlib import Path

os.environ["CONVERSION_CLEANUP_INTERVAL_MINUTES"] = "9999"

ROOT = Path(__file__).resolve().parents[1]

import sys
sys.path.insert(0, str(ROOT))

from app import create_app
from config import Config
from extensions import db
from models import ConversionJob, Usuario
import Blueprints.handlers.conversion_handlers as conversion_handler
import Blueprints.services.convertions_services.upload_flow.job_submission as job_submission
from Blueprints.services.convertions_services.errors.conversion_errors import (
    get_user_friendly_conversion_error,
)
from Blueprints.services.convertions_services.runtime.job_queue import process_conversion_job

RESULTS = []
QA_PDF_PASSWORD = "qa-senha-123"

def main():
    test_dir = Path(tempfile.mkdtemp(prefix="boost_converter_qa_"))

    try:
        Config.SQLALCHEMY_DATABASE_URI = f"sqlite:///{test_dir / 'qa.sqlite'}"
        Config.SQLALCHEMY_ENGINE_OPTIONS = {}
        app = create_app()
        # The app has its own CSRF check and a per-minute rate limit; this script posts
        # dozens of uploads in a row, so both are turned off here.
        app.config.update(
            TESTING=True
            ,CSRF_ENABLED=False
            ,RATE_LIMIT_ENABLED=False
        )
        app.instance_path = str(test_dir / "instance")
        Path(app.instance_path).mkdir(parents=True, exist_ok=True)

        patch_job_queue(app)

        with app.app_context():
            db.drop_all()
            db.create_all()
            user = Usuario(
                 email="qa@boostconverter.local"
                ,senha=None
                ,nome="QA"
                ,plano="pro"
            )
            db.session.add(user)
            db.session.commit()
            user_id = user.id

        samples = create_samples(test_dir)

        with app.test_client() as client:
            login_as(client, user_id)
            test_tool_pages(client)
            test_single_tool_conversions(client, samples)
            test_batch_zip_download(client, samples)

        with app.test_client() as client:
            test_free_file_count_limit(client, samples)

        test_error_message_sanitizer()

        print_results()

    finally:
        shutil.rmtree(test_dir, ignore_errors=True)

def patch_job_queue(app):
    def submit_sync(flask_app, job_id, convert_function, runtime_options=None):
        process_conversion_job(app, job_id, convert_function, runtime_options)

    conversion_handler.submit_conversion_job = submit_sync
    job_submission.submit_conversion_job = submit_sync

def login_as(client, user_id):
    with client.session_transaction() as session:
        session["_user_id"] = str(user_id)
        session["_fresh"] = True

def create_samples(test_dir):
    samples_dir = test_dir / "samples"
    samples_dir.mkdir(parents=True, exist_ok=True)

    create_pdf(samples_dir / "sample.pdf")
    create_pdf_with_image(samples_dir / "sample_image.pdf", samples_dir)
    create_protected_pdf(samples_dir / "protected.pdf")
    create_pptx(samples_dir / "sample.pdf", samples_dir / "sample.pptx")
    create_docx(samples_dir / "sample.docx")
    create_xlsx(samples_dir / "sample.xlsx")
    create_image(samples_dir / "sample.jpg", "JPEG")
    create_image(samples_dir / "sample.png", "PNG")
    create_image(samples_dir / "sample.webp", "WEBP")
    create_svg(samples_dir / "sample.svg")
    create_wav(samples_dir / "sample.wav")
    create_heic(samples_dir / "sample.heic")
    create_media_samples(samples_dir)

    write_text(samples_dir / "sample.csv", "name,score\nBoost,10\n")
    write_text(samples_dir / "sample.txt", "Boost Converter\nTeste rapido.\n")
    write_text(samples_dir / "sample.html", "<html><body><h1>Boost</h1><p>Teste</p></body></html>")
    write_text(samples_dir / "sample.json", '[{"name": "Boost", "score": 10}]')
    write_text(samples_dir / "sample.md", "# Boost\n\nTeste rapido.")

    return {
        "pdf": samples_dir / "sample.pdf",
        "pdf_image": samples_dir / "sample_image.pdf",
        "pdf_protected": samples_dir / "protected.pdf",
        "pptx": samples_dir / "sample.pptx",
        "docx": samples_dir / "sample.docx",
        "xlsx": samples_dir / "sample.xlsx",
        "jpg": samples_dir / "sample.jpg",
        "png": samples_dir / "sample.png",
        "webp": samples_dir / "sample.webp",
        "svg": samples_dir / "sample.svg",
        "wav": samples_dir / "sample.wav",
        "flac": samples_dir / "sample.flac",
        "mp3": samples_dir / "sample.mp3",
        "ogg": samples_dir / "sample.ogg",
        "aac": samples_dir / "sample.aac",
        "wma": samples_dir / "sample.wma",
        "mp4": samples_dir / "sample.mp4",
        "mkv": samples_dir / "sample.mkv",
        "mov": samples_dir / "sample.mov",
        "avi": samples_dir / "sample.avi",
        "webm": samples_dir / "sample.webm",
        "heic": samples_dir / "sample.heic",
        "csv": samples_dir / "sample.csv",
        "txt": samples_dir / "sample.txt",
        "html": samples_dir / "sample.html",
        "json": samples_dir / "sample.json",
        "md": samples_dir / "sample.md",
    }

def create_pdf(path):
    import pymupdf

    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Boost Converter QA")
    page.insert_text((72, 100), "Item Valor")
    page.insert_text((72, 125), "A 10")
    document.save(path)
    document.close()

def create_pdf_with_image(path, samples_dir):
    import pymupdf

    image_path = samples_dir / "embedded.png"
    create_image(image_path, "PNG")
    document = pymupdf.open()
    page = document.new_page()
    page.insert_image(pymupdf.Rect(72, 72, 272, 272), filename=str(image_path))
    document.save(path)
    document.close()

def create_protected_pdf(path):
    import pymupdf

    document = pymupdf.open()
    document.new_page().insert_text((72, 72), "Boost Converter QA protegido")
    document.save(path, encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw=QA_PDF_PASSWORD, user_pw=QA_PDF_PASSWORD)
    document.close()

def create_pptx(pdf_path, path):
    from Blueprints.services.convertions_services.documents.pdf_to_pptx_service import convert_pdf_pptx

    convert_pdf_pptx(pdf_path, path)

def create_docx(path):
    from docx import Document

    document = Document()
    document.add_heading("Boost Converter QA", 1)
    document.add_paragraph("Documento pequeno para teste.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Valor"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "10"
    document.save(path)

def create_xlsx(path):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Item", "Valor"])
    sheet.append(["A", 10])
    workbook.save(path)

def create_image(path, image_format):
    from PIL import Image

    image = Image.new("RGB", (64, 64), (30, 120, 200))
    image.save(path, image_format)

def create_svg(path):
    write_text(
        path,
        '<svg xmlns="http://www.w3.org/2000/svg" width="64" height="64">'
        '<rect width="64" height="64" fill="#1e78c8"/>'
        '<text x="8" y="36" fill="white">BC</text>'
        "</svg>",
    )

def create_wav(path):
    with wave.open(str(path), "wb") as audio:
        audio.setnchannels(1)
        audio.setsampwidth(2)
        audio.setframerate(8000)
        audio.writeframes(b"\x00\x00" * 8000)

def create_heic(path):
    try:
        import pillow_heif
        from PIL import Image

        pillow_heif.register_heif_opener()
        image = Image.new("RGB", (64, 64), (40, 160, 90))
        image.save(path, format="HEIF")
    except Exception:
        path.write_bytes(b"")

def create_media_samples(samples_dir):
    import imageio_ffmpeg

    ffmpeg = imageio_ffmpeg.get_ffmpeg_exe()
    wav_path = samples_dir / "sample.wav"

    audio_outputs = {
        "sample.flac": ["-c:a", "flac"],
        "sample.mp3": ["-c:a", "libmp3lame", "-b:a", "128k"],
        "sample.ogg": ["-c:a", "libvorbis"],
        "sample.aac": ["-c:a", "aac", "-f", "adts"],
        "sample.wma": ["-c:a", "wmav2"],
    }

    for filename, args in audio_outputs.items():
        output = samples_dir / filename
        run_ffmpeg([ffmpeg, "-y", "-i", str(wav_path), *args, str(output)])

    video_base = [
        ffmpeg,
        "-y",
        "-f", "lavfi",
        "-i", "testsrc=size=96x96:rate=10",
        "-f", "lavfi",
        "-i", "sine=frequency=1000:sample_rate=44100",
        "-t", "1",
        "-pix_fmt", "yuv420p",
    ]

    video_outputs = {
        "sample.mp4": ["-c:v", "libx264", "-c:a", "aac"],
        "sample.mkv": ["-c:v", "libx264", "-c:a", "aac"],
        "sample.mov": ["-c:v", "libx264", "-c:a", "aac"],
        "sample.avi": ["-c:v", "mpeg4", "-c:a", "mp3"],
        "sample.webm": ["-c:v", "libvpx-vp9", "-c:a", "libopus"],
    }

    for filename, args in video_outputs.items():
        output = samples_dir / filename
        run_ffmpeg([*video_base, *args, str(output)])

def run_ffmpeg(command):
    subprocess.run(command, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=True)

def write_text(path, content):
    path.write_text(content, encoding="utf-8")

def test_tool_pages(client):
    routes = [
        "/tools/aac-to-mp3",
        "/tools/pdf-merge",
        "/tools/txt-to-docx",
        "/tools/pdf-to-pptx",
        "/tools/images-to-pdf",
        "/tools/files-to-zip",
        "/tools/png-to-jpg",
        "/tools/mp4-to-gif",
    ]

    for route in routes:
        response = client.get(route)
        record(route, "tool_page", response.status_code == 200, f"status={response.status_code}")

def test_single_tool_conversions(client, samples):
    tools = {
        "/convert/aac-to-mp3": samples["aac"],
        "/convert/csv-to-xlsx": samples["csv"],
        "/convert/docx-to-pdf": samples["docx"],
        "/convert/docx-to-txt": samples["docx"],
        "/convert/docx-to-xlsx": samples["docx"],
        "/convert/html-to-docx": samples["html"],
        "/convert/html-to-pdf": samples["html"],
        "/convert/jpg-to-pdf": samples["jpg"],
        "/convert/json-to-csv": samples["json"],
        "/convert/json-to-xlsx": samples["json"],
        "/convert/md-to-docx": samples["md"],
        "/convert/md-to-pdf": samples["md"],
        "/convert/pdf-compress": samples["pdf"],
        "/convert/pdf-extract-images": samples["pdf_image"],
        "/convert/pdf-protect": (samples["pdf"], {"pdf_password": QA_PDF_PASSWORD}),
        "/convert/pdf-rotate": (samples["pdf"], {"rotation_angle": "90"}),
        "/convert/pdf-unlock": (samples["pdf_protected"], {"pdf_password": QA_PDF_PASSWORD}),
        "/convert/pptx-to-pdf": samples["pptx"],
        "/convert/pdf-edit": (samples["pdf"], {"edit_text": "QA"}),
        "/convert/pdf-merge": [samples["pdf"], samples["pdf"]],
        "/convert/pdf-split": samples["pdf"],
        "/convert/pdf-to-csv": samples["pdf"],
        "/convert/pdf-to-docx": samples["pdf"],
        "/convert/pdf-to-html": samples["pdf"],
        "/convert/pdf-to-jpg": samples["pdf"],
        "/convert/pdf-to-png": samples["pdf"],
        "/convert/pdf-to-pptx": samples["pdf"],
        "/convert/pdf-to-txt": samples["pdf"],
        "/convert/pdf-to-xlsx": samples["pdf"],
        "/convert/files-to-zip": [samples["txt"], samples["pdf"]],
        "/convert/txt-to-docx": samples["txt"],
        "/convert/txt-to-pdf": samples["txt"],
        "/convert/xlsx-to-csv": samples["xlsx"],
        "/convert/xlsx-to-docx": samples["xlsx"],
        "/convert/xlsx-to-json": samples["xlsx"],
        "/convert/xlsx-to-pdf": samples["xlsx"],
        "/convert/flac-to-mp3": samples["flac"],
        "/convert/flac-to-wav": samples["flac"],
        "/convert/heic-to-jpg": samples["heic"],
        "/convert/heic-to-png": samples["heic"],
        "/convert/images-to-pdf": [samples["jpg"], samples["png"]],
        "/convert/jpg-to-png": samples["jpg"],
        "/convert/jpg-to-svg": samples["jpg"],
        "/convert/jpg-to-webp": samples["jpg"],
        "/convert/mp3-to-mp4": samples["mp3"],
        "/convert/mp3-to-wav": samples["mp3"],
        "/convert/ogg-to-mp3": samples["ogg"],
        "/convert/ogg-to-wav": samples["ogg"],
        "/convert/png-to-jpg": samples["png"],
        "/convert/png-to-svg": samples["png"],
        "/convert/png-to-webp": samples["png"],
        "/convert/svg-to-jpg": samples["svg"],
        "/convert/svg-to-png": samples["svg"],
        "/convert/webp-to-jpg": samples["webp"],
        "/convert/webp-to-png": samples["webp"],
        "/convert/wav-to-flac": samples["wav"],
        "/convert/wav-to-mp3": samples["wav"],
        "/convert/wma-to-mp3": samples["wma"],
        "/convert/avi-to-mp4": samples["avi"],
        "/convert/mkv-to-mp4": samples["mkv"],
        "/convert/mov-to-mp4": samples["mov"],
        "/convert/mp4-to-gif": samples["mp4"],
        "/convert/mp4-to-mkv": samples["mp4"],
        "/convert/mp4-to-mov": samples["mp4"],
        "/convert/mp4-to-mp3": samples["mp4"],
        "/convert/mp4-to-wav": samples["mp4"],
        "/convert/mp4-to-webm": samples["mp4"],
        "/convert/webm-to-mp4": samples["webm"],
    }

    for route, config in tools.items():
        form = None
        path = config

        if isinstance(config, tuple):
            path, form = config

        paths = path if isinstance(path, list) else [path]
        run_conversion(client, route, paths, form=form)

def test_batch_zip_download(client, samples):
    response = post_files(client, "/convert/txt-to-docx", [samples["txt"], samples["txt"]])

    if response.status_code not in (302, 303):
        record("/convert/txt-to-docx", "batch_upload", False, f"status={response.status_code}")
        return

    location = response.headers.get("Location", "")
    batch_id = location.rsplit("/", 1)[-1]
    response = client.get(f"/conversions/batch/{batch_id}/download")
    record(
        "/conversions/batch/download",
        "zip_download",
        response.status_code == 200 and response.data.startswith(b"PK"),
        f"status={response.status_code}, bytes={len(response.data)}",
    )

def test_free_file_count_limit(client, samples):
    ok_response = post_files(client, "/convert/txt-to-docx", [samples["txt"], samples["txt"]])
    record(
        "/convert/txt-to-docx",
        "free_two_files",
        ok_response.status_code in (302, 303),
        f"status={ok_response.status_code}",
    )

    upgrade_response = post_files(
        client,
        "/convert/txt-to-docx",
        [samples["txt"], samples["txt"], samples["txt"]],
    )
    record(
        "/convert/txt-to-docx",
        "free_three_files_requires_upgrade",
        is_redirect_to_plans(upgrade_response),
        f"status={upgrade_response.status_code}, location={upgrade_response.headers.get('Location', '')}",
    )

    merge_ok = post_files(client, "/convert/pdf-merge", [samples["pdf"]] * 32)
    record("/convert/pdf-merge", "free_32_pdfs", merge_ok.status_code in (302, 303) and not is_redirect_to_plans(merge_ok),
           f"status={merge_ok.status_code}, location={merge_ok.headers.get('Location', '')}")
    merge_over = post_files(client, "/convert/pdf-merge", [samples["pdf"]] * 33)
    record("/convert/pdf-merge", "free_33_pdfs_requires_upgrade", is_redirect_to_plans(merge_over),
           f"status={merge_over.status_code}, location={merge_over.headers.get('Location', '')}")

def is_redirect_to_plans(response):
    return response.status_code in (302, 303) and response.headers.get("Location", "").endswith("/planos")

def test_error_message_sanitizer():
    technical_error = FileNotFoundError(
        "[Errno 2] No such file or directory: "
        "C:\\Users\\teste\\arquivo.mp4"
    )
    friendly_message = get_user_friendly_conversion_error(technical_error)

    record(
        "conversion_errors",
        "friendly_file_not_found",
        "C:\\" not in friendly_message and "Errno" not in friendly_message,
        friendly_message,
    )

    validation_error = ValueError("Informe um texto para adicionar ao PDF.")
    friendly_message = get_user_friendly_conversion_error(validation_error)

    record(
        "conversion_errors",
        "friendly_validation_error",
        friendly_message == "Informe um texto para adicionar ao PDF.",
        friendly_message,
    )

def run_conversion(client, route, paths, form=None):
    response = post_files(client, route, paths, form=form)

    if response.status_code not in (302, 303):
        record(route, "upload", False, f"status={response.status_code}, body={response.data[:180]!r}")
        return

    job = get_last_job()

    if job is None:
        record(route, "job", False, "job nao encontrado")
        return

    ok = job.status == "done" and os.path.exists(job.output_path) and os.path.getsize(job.output_path) > 0
    detail = f"job={job.status}, output={job.output_filename}, error={job.error_message}"
    record(route, "conversion", ok, detail)

    if ok:
        format_ok, format_detail = validate_output(job.output_path, job.output_filename)
        record(route, "output_format", format_ok, format_detail)

    if ok:
        response = client.get(f"/conversions/{job.id}/download")
        record(route, "download", response.status_code == 200 and len(response.data) > 0, f"status={response.status_code}")

BINARY_SIGNATURES = {
    "pdf": [(0, b"%PDF-")],
    "png": [(0, b"\x89PNG")],
    "jpg": [(0, b"\xff\xd8\xff")],
    "gif": [(0, b"GIF8")],
    "webp": [(0, b"RIFF"), (8, b"WEBP")],
    "zip": [(0, b"PK")],
    "docx": [(0, b"PK")],
    "xlsx": [(0, b"PK")],
    "pptx": [(0, b"PK")],
    "wav": [(0, b"RIFF"), (8, b"WAVE")],
    "avi": [(0, b"RIFF"), (8, b"AVI ")],
    "flac": [(0, b"fLaC")],
    "mp4": [(4, b"ftyp")],
    "mov": [(4, b"ftyp")],
    "mkv": [(0, b"\x1a\x45\xdf\xa3")],
    "webm": [(0, b"\x1a\x45\xdf\xa3")],
}

def validate_output(output_path, output_filename):
    """Checks that the file really is the promised format, not just that it is not empty."""
    extension = output_filename.rsplit(".", 1)[-1].lower()
    data = Path(output_path).read_bytes()
    if extension in BINARY_SIGNATURES:
        ok = all(data[offset:offset + len(signature)] == signature for offset, signature in BINARY_SIGNATURES[extension])
        return ok, f"{extension} signature={'ok' if ok else data[:12]!r}"
    if extension == "mp3":
        ok = data[:3] == b"ID3" or (data[0] == 0xFF and data[1] & 0xE0 == 0xE0)
        return ok, f"mp3 header={data[:4]!r}"
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return False, f"{extension} is not UTF-8 text"
    if extension == "json":
        import json

        try:
            json.loads(text)
        except ValueError as error:
            return False, f"invalid json: {error}"
        return True, "json parses"
    if extension == "svg":
        return "<svg" in text, "svg root element"
    if extension in {"html", "htm"}:
        return "<html" in text.lower() or "<body" in text.lower(), "html document"
    return bool(text.strip()), f"{extension} has text"

def post_files(client, route, paths, form=None):
    upload_files = []

    for index, path in enumerate(paths, start=1):
        upload_files.append((open(path, "rb"), f"{index}_{path.name}"))

    data = {"files": upload_files}

    if form:
        data.update(form)

    try:
        return client.post(
            route,
            data=data,
            content_type="multipart/form-data",
            follow_redirects=False,
        )
    finally:
        for file_object, _filename in upload_files:
            file_object.close()

def get_last_job():
    return ConversionJob.query.order_by(ConversionJob.created_at.desc()).first()

def record(route, check, ok, detail):
    status = "PASS" if ok else "FAIL"
    RESULTS.append((status, route, check, detail))

def print_results():
    for status, route, check, detail in RESULTS:
        print(f"{status}\t{check}\t{route}\t{detail}")

    failures = [result for result in RESULTS if result[0] == "FAIL"]
    print(f"\nSUMMARY total={len(RESULTS)} failed={len(failures)}")

    if failures:
        raise SystemExit(1)

if __name__ == "__main__":
    main()
