from docx import Document

def convert_docx_txt(input_path, output_path):
    document = Document(input_path)
    text_parts = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip(): text_parts.append(paragraph.text)

    for tabela_index, table in enumerate(document.tables, start=1):
        text_parts.append(f"\n--- Tabela {tabela_index} ---")

        for row in table.rows:
            row_text = []

            for cell in row.cells: row_text.append(cell.text.strip())

            text_parts.append(" | ".join(row_text))

    full_text = "\n".join(text_parts)

    with open(output_path, "w", encoding="utf-8") as file:
        file.write(full_text)