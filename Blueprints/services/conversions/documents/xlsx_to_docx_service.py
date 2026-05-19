import pandas
from docx import Document

def convert_excel_docx(input_path, ouput_path):
    dataframe = pandas.read_excel(input_path)
    doc = Document()
    table = doc.add_table(rows=dataframe.shape[0] + 1, cols=dataframe.shape[1])

    for col_index, col_name in enumerate(dataframe.columns):
        table.cell(0, col_index).text = str(col_name)

    for row_index, row in dataframe.iterrows():
        for col_index, value in enumerate(row):
            table.cell(row_index + 1, col_index).text = str(value)

    doc.save(ouput_path)