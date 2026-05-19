from docx import Document

def convert_markdown_docx(input_path, output_path):
    document = Document()

    with open(input_path, "r", encoding="utf-8") as file:
        for raw_line in file:
            line = raw_line.strip()

            if not line:
                continue

            if line.startswith("# "):
                document.add_heading(line[2:].strip(), level=1)
            elif line.startswith("## "):
                document.add_heading(line[3:].strip(), level=2)
            elif line.startswith("### "):
                document.add_heading(line[4:].strip(), level=3)
            else:
                document.add_paragraph(clean_markdown_line(line))

    document.save(output_path)

def clean_markdown_line(line):
    return line.replace("**", "").replace("__", "").replace("*", "").replace("`", "")
